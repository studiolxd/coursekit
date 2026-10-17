"""`coursekit design CODE`: the instructional design of a course made in the project, without slxd.

A course whose `course.yaml › design.tool` is `local` keeps its design in `courses/<CODE>/design/design.yaml` (written by
the design agent or by a person; English keys, texts in the course language). This module checks it with the rules of a
design matrix (coverage, evidence, practice before assessment, instrument and Bloom level, hours, verbs, progression) and
exports what the rest of coursekit reads, the same three files a design made in slxd leaves:

- `design/matrix.json`: the graph in the shape of slxd's `design_matrix_get` (so `sync`, `approve` and the assembly do not
  care where the design was made), with ids derived from the course code and the ids of the file, stable between exports;
- `design/validation.json`: the findings, in the shape of `design_matrix_validate`;
- `design/<CODE>-instructional-design.xlsx`: what the person reviews.
"""

from __future__ import annotations

import json
import re
import unicodedata
import uuid
from dataclasses import dataclass
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from coursekit import course as coursemod
from coursekit import lang
from coursekit.i18n import t

FILE = Path("design") / "design.yaml"
NAMESPACE = uuid.UUID("6f1c2a52-6f0e-4d4b-9a59-0c2a3b7d4e10")

BLOOM = ("remember", "understand", "apply", "analyze", "evaluate", "create")
BLOOM_SLXD = dict(zip(BLOOM, ("recordar", "comprender", "aplicar", "analizar", "evaluar", "crear"), strict=True))
LEARNING_TYPES = {"self_assessed": "autoevaluable", "classification": "clasificacion", "interactive_decision": "interactiva_decision",
                  "branching_scenario": "escenario_ramificado", "hotspot_analysis": "analisis_hotspots",
                  "guided_simulation": "simulacion_guiada", "template_practice": "practica_plantilla"}
INSTRUMENTS = {"quiz": "cuestionario", "case_study": "caso_practico", "portfolio": "entrega_portfolio", "rubric": "rubrica"}
CEILING = {"quiz": "apply", "case_study": "evaluate", "portfolio": "create", "rubric": "create"}  # highest level each one evidences
PASSING = {"all_tests": "todos_los_examenes", "course_score": "nota_del_curso"}
STRUCTURES = ("sin_modulos", "con_modulos")
ID_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_.-]*$")


class DesignError(Exception):
    pass


@dataclass
class Finding:
    check: str
    severity: str  # error | warning
    message: str
    kind: str = "matriz"  # what it points at, as slxd names it: competencia | objetivo | nodo | actividad | matriz
    ref: str | None = None

    def line(self) -> str:
        mark = "✘" if self.severity == "error" else "·"
        return f"{mark} {self.message}"


def is_local(course: dict) -> bool:
    """The course designs in the project (`design.tool: local`); without the key, in slxd."""
    return str((course.get("design") or {}).get("tool") or "slxd") == "local"


def path(course: dict) -> Path:
    return course["_dir"] / FILE


def load(course: dict) -> dict:
    file = path(course)
    if not file.is_file():
        raise DesignError(t("design", "missing", path=FILE.as_posix()))
    try:
        data = yaml.safe_load(file.read_text(encoding="utf-8"))
    except yaml.YAMLError as exc:
        raise DesignError(t("design", "bad_yaml", path=FILE.as_posix(), error=str(exc).splitlines()[0])) from exc
    if not isinstance(data, dict):
        raise DesignError(t("design", "not_a_mapping", path=FILE.as_posix()))
    return data


def _id(code: str, *parts: object) -> str:
    return str(uuid.uuid5(NAMESPACE, "/".join([code, *map(str, parts)])))


def _list(value: object) -> list:
    return value if isinstance(value, list) else []


def _hours(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return round(float(value), 2)


def _number(value: float) -> float | int:
    value = round(float(value), 2)
    return int(value) if value.is_integer() else value


def _plain(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", text.casefold()) if unicodedata.category(c) != "Mn")


def first_verb(text: str, language: str) -> str:
    """The verb an objective starts with, lower case and without accents ("Aplicar los criterios…" → "aplicar")."""
    words = re.findall(r"[^\W\d_]+", text)
    if language == "en" and words and words[0].casefold() == "to":
        words = words[1:]
    return _plain(words[0]) if words else ""


class Builder:
    """design.yaml → (graph, findings)."""

    def __init__(self, course: dict, data: dict):
        self.course = course
        self.data = data
        self.code = str(course["code"])
        self.language = coursemod.language(course)
        self.findings: list[Finding] = []
        self.unit_word = coursemod.tokens(course)["unit"]

    def add(self, check: str, severity: str, key: str, kind: str = "matriz", ref: str | None = None, **values) -> None:
        self.findings.append(Finding(check, severity, t("design", key, **values), kind, ref))

    # ── reading ─────────────────────────────────────────────────────────────────────────

    def build(self) -> dict:
        design = self.course.get("design") or {}
        profile = self.data.get("profile") if isinstance(self.data.get("profile"), dict) else {}
        audience = coursemod.audience(self.course)  # the project's, changed by the course
        structure = str(design.get("structure") or "sin_modulos")
        cfg = coursemod.rules(self.course)
        passing = str(self.data.get("passing") or "all_tests")
        if passing not in PASSING:
            self.add("structure.invalid", "error", "bad_value", field="passing", value=passing, valid=", ".join(PASSING))
            passing = "all_tests"
        course_score = self.data.get("course_passing_score")
        if passing == "course_score" and not isinstance(course_score, int):
            self.add("structure.invalid", "error", "course_score_needed")
        self.graph: dict = {
            "matrix": {
                "id": _id(self.code, "matrix"),
                "nombre": self.course.get("title"),
                "duracionHoras": _number(float(design.get("hours") or 0)),
                "nivel": profile.get("level") or audience["level"] or None,
                "publicoObjetivo": profile.get("audience") or audience["who"] or None,
                "conocimientosPrevios": profile.get("prerequisites") or audience["knowledge"] or None,
                "contexto": profile.get("context") or audience["context"] or None,
                "tono": profile.get("tone") or design.get("tone") or None,
                "patronEstructural": structure,
                "palabrasPorHora": int(cfg["pages_per_hour"]) * int(cfg["words_per_page"]),
                "criterioAprobado": PASSING[passing],
                "notaAprobadoCurso": course_score if passing == "course_score" else None,
                "origen": "local",
                "updatedAt": None,
            },
            "competencias": [],
            "objetivos": [],
            "nodos": [],
            "actividadesAprendizaje": [],
            "actividadesEvaluacion": [],
        }
        self.competencies = self._competencies()
        self.objectives = self._objectives()
        self.links: dict[str, set[str]] = {key: set() for key in self.objectives}  # objective → leaf nodes
        self.practised: dict[str, set[str]] = {key: set() for key in self.objectives}  # objective → units with a learning activity
        self.assessed: dict[str, list[str]] = {key: [] for key in self.objectives}  # objective → instruments of its evidence
        self.unit_objectives: list[tuple[str, str, list[str]]] = []  # (node id, title, objective ids in its sections)
        modules = _list(self.data.get("modules"))
        units = _list(self.data.get("units"))
        if structure == "con_modulos":
            if units:
                self.add("structure.invalid", "error", "units_need_modules")
            if not modules:
                self.add("structure.empty", "error", "no_modules")
            for mn, module in enumerate(modules, 1):
                self._module(module, mn)
        else:
            if modules:
                self.add("structure.invalid", "error", "modules_not_allowed")
            if not units:
                self.add("structure.empty", "error", "no_units")
            for un, unit in enumerate(units, 1):
                self._unit(unit, un, None, f"u{un}")
        self._checks()
        return self.graph

    def _competencies(self) -> dict[str, str]:
        found: dict[str, str] = {}
        for order, item in enumerate(_list(self.data.get("competencies"))):
            if not isinstance(item, dict) or not str(item.get("text") or "").strip():
                self.add("structure.invalid", "error", "competency_without_text", n=order + 1)
                continue
            key = str(item.get("id") or f"C{order + 1}")
            if key in found:
                self.add("structure.duplicate", "error", "duplicate_id", id=key)
                continue
            found[key] = _id(self.code, "competency", key)
            self.graph["competencias"].append({"id": found[key], "texto": str(item["text"]).strip(), "orden": order, "clave": key})
        return found

    def _objectives(self) -> dict[str, dict]:
        found: dict[str, dict] = {}
        for order, item in enumerate(_list(self.data.get("objectives"))):
            if not isinstance(item, dict):
                self.add("structure.invalid", "error", "objective_not_mapping", n=order + 1)
                continue
            key = str(item.get("id") or "")
            if not ID_RE.match(key):
                self.add("structure.invalid", "error", "objective_bad_id", n=order + 1, id=key)
                continue
            if key in found or key in self.competencies:
                self.add("structure.duplicate", "error", "duplicate_id", id=key)
                continue
            text = str(item.get("text") or "").strip()
            bloom = str(item.get("bloom") or "")
            if not text:
                self.add("structure.invalid", "error", "objective_without_text", id=key)
            if bloom not in BLOOM:
                self.add("structure.invalid", "error", "bad_value", field=f"{key}.bloom", value=bloom, valid=", ".join(BLOOM))
                bloom = "apply"
            competencies = []
            for ref in _list(item.get("competencies")):
                if str(ref) in self.competencies:
                    competencies.append(self.competencies[str(ref)])
                else:
                    self.add("structure.reference", "error", "unknown_competency", id=key, name=ref)
            node = {"id": _id(self.code, "objective", key), "texto": text, "nivelBloom": BLOOM_SLXD[bloom], "orden": order,
                    "competenciaIds": competencies, "clave": key}
            found[key] = {"node": node, "bloom": bloom, "text": text, "order": order}
            self.graph["objetivos"].append(node)
        return found

    def _objective_refs(self, value: object, where: str) -> list[str]:
        out = []
        for ref in _list(value):
            if str(ref) in self.objectives:
                if str(ref) not in out:
                    out.append(str(ref))
            else:
                self.add("structure.reference", "error", "unknown_objective", where=where, name=ref)
        return out

    def _node(self, key: str, parent: str | None, kind: str, title: str, order: int, hours: float | None,
              objectives: list[str], general: str | None = None) -> dict:
        node = {"id": _id(self.code, "node", key), "parentId": parent, "tipo": kind, "titulo": title, "orden": order,
                "horas": hours, "objetivoGeneral": general,
                "objetivoIds": [self.objectives[o]["node"]["id"] for o in objectives]}
        self.graph["nodos"].append(node)
        return node

    def _module(self, module: object, mn: int) -> None:
        if not isinstance(module, dict) or not str(module.get("title") or "").strip():
            self.add("structure.invalid", "error", "module_without_title", n=mn)
            return
        key = str(module.get("id") or f"m{mn}")
        node = self._node(key, None, "modulo", str(module["title"]).strip(), mn - 1, _hours(module.get("hours")), [],
                          str(module.get("general_objective") or "").strip() or None)
        units = _list(module.get("units"))
        if not units:
            self.add("structure.empty", "error", "module_without_units", title=node["titulo"])
        children = [self._unit(unit, un, node["id"], f"{key}/u{un}") for un, unit in enumerate(units, 1)]
        self._sum(node, [c for c in children if c])

    def _unit(self, unit: object, un: int, parent: str | None, default_key: str) -> dict | None:
        n = len(self.unit_objectives) + 1
        if not isinstance(unit, dict) or not str(unit.get("title") or "").strip():
            self.add("structure.invalid", "error", "unit_without_title", unit=self.unit_word, n=n)
            return None
        key = str(unit.get("id") or default_key)
        title = str(unit["title"]).strip()
        node = self._node(key, parent, "unidad", title, un - 1, _hours(unit.get("hours")), [])
        sections = _list(unit.get("sections"))
        if not sections:
            self.add("structure.empty", "error", "unit_without_sections", unit=self.unit_word, n=n, kind="nodo", ref=node["id"])
        design = self.course.get("design") or {}
        intro, summary = bool(design.get("intro_section", True)), bool(design.get("summary_section", True))
        children, covered = [], []
        for sn, section in enumerate(sections, 1):
            framing = (intro and sn == 1) or (summary and sn == len(sections))  # introduction and summary develop no objective
            child = self._section(section, sn, node["id"], f"{key}/s{sn}", "apartado", n, framing)
            if child:
                children.append(child[0])
                covered += [o for o in child[1] if o not in covered]
        self._sum(node, children)
        self.unit_objectives.append((node["id"], title, covered))
        for order, item in enumerate(_list(unit.get("learning_activities"))):
            self._activity(item, order, node["id"], f"{key}/la{order + 1}", n, learning=True)
        for order, item in enumerate(_list(unit.get("assessment_activities"))):
            self._activity(item, order, node["id"], f"{key}/aa{order + 1}", n, learning=False)
        return node

    def _section(self, section: object, sn: int, parent: str, key: str, kind: str, un: int,
                 framing: bool) -> tuple[dict, list[str]] | None:
        if not isinstance(section, dict) or not str(section.get("title") or "").strip():
            self.add("structure.invalid", "error", "section_without_title", unit=self.unit_word, n=un, s=sn)
            return None
        if section.get("id"):  # an id of its own keeps the identity of the section when the ones before it move
            key = f"{key.rsplit('/', 1)[0]}/{section['id']}"
        title = str(section["title"]).strip()
        subsections = _list(section.get("subsections")) if kind == "apartado" else []
        objectives = self._objective_refs(section.get("objectives"), title)
        if subsections and objectives:
            self.add("structure.invalid", "error", "objectives_on_parent", title=title)
        node = self._node(key, parent, kind, title, sn - 1, _hours(section.get("hours")), [] if subsections else objectives)
        if not subsections:
            for o in objectives:
                self.links[o].add(node["id"])
            if not objectives and not framing:
                self.add("coverage.orphan_content", "warning", "orphan_content", title=title, kind="nodo", ref=node["id"])
            if node["horas"] is None:
                self.add("hours.missing", "error", "section_without_hours", title=title, kind="nodo", ref=node["id"])
            return node, objectives
        covered: list[str] = []
        children = []
        for ssn, sub in enumerate(subsections, 1):
            child = self._section(sub, ssn, node["id"], f"{key}/{ssn}", "subapartado", un, framing)
            if child:
                children.append(child[0])
                covered += [o for o in child[1] if o not in covered]
        self._sum(node, children)
        return node, covered

    def _sum(self, node: dict, children: list[dict]) -> None:
        """A node without hours takes its children's; with hours, they must add up to them."""
        hours = [c.get("horas") for c in children]
        if not children or any(h is None for h in hours):
            return
        total = round(sum(hours), 2)
        if node["horas"] is None:
            node["horas"] = _number(total)
        elif abs(node["horas"] - total) > 0.01:
            self.add("hours.children", "error", "hours_children", title=node["titulo"], declared=_number(node["horas"]),
                     children=_number(total), kind="nodo", ref=node["id"])

    def _activity(self, item: object, order: int, unit: str, key: str, un: int, learning: bool) -> None:
        which = "learning_activities" if learning else "assessment_activities"
        if not isinstance(item, dict) or not str(item.get("title") or "").strip():
            self.add("structure.invalid", "error", "activity_without_title", unit=self.unit_word, n=un, which=which, i=order + 1)
            return
        title = str(item["title"]).strip()
        objectives = self._objective_refs(item.get("objectives"), title)
        base = {"id": _id(self.code, "activity", key), "nodeId": unit, "titulo": title,
                "descripcion": str(item.get("description") or "").strip() or None, "orden": order,
                "objetivoIds": [self.objectives[o]["node"]["id"] for o in objectives]}
        if learning:
            kind = str(item.get("type") or "")
            if kind not in LEARNING_TYPES:
                self.add("structure.invalid", "error", "bad_value", field=f"{title} › type", value=kind, valid=", ".join(LEARNING_TYPES))
                kind = "self_assessed"
            self.graph["actividadesAprendizaje"].append({**base, "tipo": LEARNING_TYPES[kind]})
            for o in objectives:
                self.practised[o].add(unit)
            return
        instrument = str(item.get("instrument") or "")
        if instrument not in INSTRUMENTS:
            self.add("structure.invalid", "error", "bad_value", field=f"{title} › instrument", value=instrument,
                     valid=", ".join(INSTRUMENTS))
            instrument = "quiz"
        grading = {}
        for name, slxd, top in (("weight", "peso", 1000), ("passing_score", "notaAprobado", 100), ("attempts", "intentosMax", 100)):
            value = item.get(name)
            if value is not None and (isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= top):
                self.add("structure.invalid", "error", "bad_number", field=f"{title} › {name}", value=value, top=top)
                value = None
            grading[slxd] = value
        self.graph["actividadesEvaluacion"].append({**base, "instrumento": INSTRUMENTS[instrument],
                                                    "criterio": str(item.get("criterion") or "").strip() or None, **grading})
        for o in objectives:
            self.assessed[o].append(instrument)

    # ── the checks of a design matrix ───────────────────────────────────────────────────

    def _checks(self) -> None:
        verbs = lang.tokens(self.language).get("bloom_verbs") or {}
        levels = {_plain(str(v)): level for level in BLOOM for v in verbs.get(level) or []}
        vague = {_plain(str(v)) for v in lang.tokens(self.language).get("vague_verbs") or []}
        used = {c for o in self.objectives.values() for c in o["node"]["competenciaIds"]}
        for competency in self.graph["competencias"]:
            if competency["id"] not in used:
                self.add("coverage.competency_without_objective", "warning", "competency_without_objective",
                         id=competency["clave"], kind="competencia", ref=competency["id"])
        for key, objective in self.objectives.items():
            ref = objective["node"]["id"]
            if self.competencies and not objective["node"]["competenciaIds"]:
                self.add("coverage.objective_without_competency", "warning", "objective_without_competency", id=key,
                         kind="objetivo", ref=ref)
            if not self.links[key]:
                self.add("coverage.objective_without_content", "error", "objective_without_content", id=key, kind="objetivo", ref=ref)
            if not self.assessed[key]:
                self.add("evidence.objective_without_assessment", "error", "objective_without_assessment", id=key,
                         kind="objetivo", ref=ref)
            elif not self.practised[key]:
                self.add("arc.assessed_without_practice", "warning", "assessed_without_practice", id=key, kind="objetivo", ref=ref)
            level = BLOOM.index(objective["bloom"])
            if self.assessed[key] and all(BLOOM.index(CEILING[i]) < level for i in self.assessed[key]):
                self.add("rigor.instrument_ceiling", "error", "instrument_ceiling", id=key, bloom=objective["bloom"],
                         instruments=", ".join(sorted(set(self.assessed[key]))), kind="objetivo", ref=ref)
            verb = first_verb(objective["text"], self.language)
            if verb in vague:
                self.add("verbs.vague", "warning", "vague_verb", id=key, verb=verb, kind="objetivo", ref=ref)
            elif verb in levels and levels[verb] != objective["bloom"]:
                self.add("verbs.bloom_mismatch", "warning", "bloom_mismatch", id=key, verb=verb, level=levels[verb],
                         bloom=objective["bloom"], kind="objetivo", ref=ref)
        reached = -1  # the highest Bloom level of the earlier units
        for node, title, keys in self.unit_objectives:
            top = max((BLOOM.index(self.objectives[k]["bloom"]) for k in keys), default=-1)
            if top >= BLOOM.index("evaluate") and reached < BLOOM.index("apply"):
                self.add("progression.jump_without_base", "warning", "jump_without_base", title=title, kind="nodo", ref=node)
            reached = max(reached, top)
            if not keys:
                self.add("coverage.unit_without_objectives", "warning", "unit_without_objectives", title=title, kind="nodo", ref=node)
        total = round(sum(n["horas"] or 0 for n in self.graph["nodos"] if n["parentId"] is None), 2)
        hours = float((self.course.get("design") or {}).get("hours") or 0)
        if self.graph["nodos"] and abs(total - hours) > 0.01:
            self.add("hours.total", "error", "hours_total", total=_number(total), hours=_number(hours))


def build(course: dict) -> tuple[dict, list[Finding]]:
    builder = Builder(course, load(course))
    graph = builder.build()
    return graph, builder.findings


def validation(findings: list[Finding]) -> dict:
    return {"hallazgos": [{"check": f.check, "severidad": "error" if f.severity == "error" else "aviso", "mensaje": f.message,
                           "ref": {"tipo": f.kind, "id": f.ref}} for f in findings]}


def matrix_text(graph: dict) -> str:
    return json.dumps(graph, ensure_ascii=False, indent=2) + "\n"


def excel_name(course: dict) -> str:
    return f"{course['code']}-instructional-design.xlsx"


def stale(course: dict) -> bool:
    """design.yaml says something else than the exported matrix.json (it was edited and not exported again)."""
    matrix = course["_dir"] / "design" / "matrix.json"
    try:
        graph, _ = build(course)
    except DesignError:
        return True
    return not matrix.exists() or matrix.read_text(encoding="utf-8") != matrix_text(graph)


def export_problem(course: dict) -> str | None:
    """Why the export of a local design cannot be used (never exported, or design.yaml changed since); None when it can, or
    when the course designs in slxd."""
    if not is_local(course):
        return None
    if not (course["_dir"] / "design" / "matrix.json").is_file():
        return t("design", "not_exported", code=course["code"])
    return t("design", "stale", code=course["code"]) if stale(course) else None


def export(course: dict) -> tuple[list[Finding], list[Path]]:
    """Checks design.yaml and writes matrix.json, validation.json and the Excel."""
    graph, findings = build(course)
    folder = course["_dir"] / "design"
    folder.mkdir(parents=True, exist_ok=True)
    matrix = folder / "matrix.json"
    matrix.write_text(matrix_text(graph), encoding="utf-8", newline="\n")
    report = folder / "validation.json"
    report.write_text(json.dumps(validation(findings), ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    excel = folder / excel_name(course)
    write_excel(course, graph, findings, excel)
    return findings, [matrix, report, excel]


# ── the Excel the person reviews ─────────────────────────────────────────────────────────

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
UNIT_FONT = Font(bold=True)


def _sheet(book: Workbook, title: str, columns: list[tuple[str, int]], first: bool = False):
    sheet = book.active if first else book.create_sheet()
    sheet.title = title[:31]
    for i, (name, width) in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=i, value=name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def _row(sheet, values: list, bold: bool = False) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if bold:
            cell.font = UNIT_FONT


def keys(course: dict, graph: dict) -> tuple[dict[str, str], dict[str, str], dict[str, str]]:
    """What a person reads for each competency (C1…), each objective (U1.1…, the ids `coursekit sync` gives and
    misconceptions.yaml uses; with the id of design.yaml after it when that one differs) and the kind of each section node
    (intro, content, activities, summary)."""
    from coursekit import sync as syncmod  # imports this module

    competencies = {c["id"]: str(c.get("clave") or f"C{n}") for n, c in enumerate(syncmod.ordered(graph["competencias"]), 1)}
    units, _ = syncmod.build_units(graph, course, coursemod.rules(course), coursemod.tokens(course))
    labels = {str(o["slxd_id"]): str(o["id"]) for u in units for o in u["objectives"]}
    objectives = {}
    for n, o in enumerate(syncmod.ordered(graph["objetivos"]), 1):
        label, own = labels.get(o["id"]), o.get("clave")
        objectives[o["id"]] = (f"{label} ({own})" if own and own != label else label) if label else str(own or f"O{n}")
    kinds = {s["slxd_id"]: s["kind"] for u in units for s in u["sections"]}
    return competencies, objectives, kinds


ORPHAN_CHECKS = ("coverage.orphan_content", "cobertura.contenidoHuerfano")


def slxd_findings(path: Path) -> list[Finding]:
    """The findings of `design_matrix_validate` that the design agent saved in validation.json."""
    if not path.is_file():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    items = data.get("hallazgos", []) if isinstance(data, dict) else data
    return [Finding(str(h.get("check") or ""), "error" if h.get("severidad") == "error" else "warning", str(h.get("mensaje") or ""),
                    str((h.get("ref") or {}).get("tipo") or "matriz"), (h.get("ref") or {}).get("id")) for h in items or []]


def review_excel(course: dict) -> tuple[list[Finding], Path]:
    """A course designed in slxd: the Excel the person reviews, the same workbook as a local design, from the matrix.json and
    validation.json the design agent saved (the Excel slxd exports lacks the competencies, the hours of the sections and the
    grading of the tests)."""
    from coursekit import sync as syncmod  # imports this module

    folder = course["_dir"] / "design"
    graph = syncmod.load_graph(folder / "matrix.json")
    findings = slxd_findings(folder / "validation.json") + key_findings(course, graph)
    target = folder / excel_name(course)
    write_excel(course, graph, findings, target)
    return findings, target


def validation_problem(course: dict, required: bool = True) -> str | None:
    """A design made in slxd: why its validation.json does not belong to its matrix.json (missing, when `required`, or saved
    before the matrix last changed: `design_matrix_validate` was not run again). None when it does, and for a local design
    (`coursekit design` writes both). slxd dates neither, so the dates of the files tell it."""
    if is_local(course):
        return None
    folder = course["_dir"] / "design"
    matrix, validation = folder / "matrix.json", folder / "validation.json"
    if not matrix.is_file():
        return None
    if not validation.is_file():
        return t("design", "validation_missing", code=course["code"]) if required else None
    if validation.stat().st_mtime < matrix.stat().st_mtime - 1:
        return t("design", "validation_stale", code=course["code"])
    return None


def key_findings(course: dict, graph: dict) -> list[Finding]:
    """A design made in slxd: the objectives whose `clave` (what slxd's Excel shows) is not the id coursekit gives them."""
    from coursekit import sync as syncmod  # imports this module

    units, _ = syncmod.build_units(graph, course, coursemod.rules(course), coursemod.tokens(course))
    labels = {str(o["slxd_id"]): str(o["id"]) for u in units for o in u["objectives"]}
    return [Finding("keys.mismatch", "warning", t("design", "key_differs", key=o["clave"], id=labels[o["id"]], text=o.get("texto")),
                    "objetivo", o["id"])
            for o in graph.get("objetivos") or [] if o.get("clave") and o["id"] in labels and o["clave"] != labels[o["id"]]]


def expected(finding: Finding, kinds: dict[str, str]) -> bool:
    """A warning the design asks for: an introduction or a summary develops no objective."""
    return finding.check in ORPHAN_CHECKS and kinds.get(str(finding.ref)) in ("intro", "summary")


def write_excel(course: dict, graph: dict, findings: list[Finding], target: Path) -> None:
    from coursekit import misconceptions as misconmod

    head = lambda key: t("design", key)  # noqa: E731
    bloom_name = {BLOOM_SLXD[b]: t("design", f"bloom_{b}") for b in BLOOM}
    kinds = {v: t("design", f"type_{k}") for k, v in LEARNING_TYPES.items()}
    instruments = {v: t("design", f"instrument_{k}") for k, v in INSTRUMENTS.items()}
    competencies, objectives, section_kinds = keys(course, graph)
    nodes = {n["id"]: n for n in graph["nodos"]}
    rules = coursemod.rules(course)
    names = lambda ids, found: ", ".join(found.get(i, "?") for i in ids)  # noqa: E731
    ordered = lambda items: sorted(items, key=lambda x: x.get("orden", 0))  # noqa: E731

    book = Workbook()
    meta = graph["matrix"]
    passing = next((k for k, v in PASSING.items() if v == meta.get("criterioAprobado")), "all_tests")
    leaves = [n for n in graph["nodos"] if n["tipo"] in ("apartado", "subapartado")
              and not any(c["parentId"] == n["id"] for c in graph["nodos"])]
    words = sum(coursemod.min_words(n.get("horas") or 0, rules) for n in leaves)
    sheet = _sheet(book, head("sheet_profile"), [(head("col_field"), 28), (head("col_value"), 90)], first=True)
    for key, value in (("field_title", meta.get("nombre")), ("field_code", course["code"]), ("field_hours", meta.get("duracionHoras")),
                       ("field_min_words", words),
                       ("field_structure", meta.get("patronEstructural")), ("field_words_per_hour", meta.get("palabrasPorHora")),
                       ("field_audience", meta.get("publicoObjetivo")), ("field_level", meta.get("nivel")),
                       ("field_prerequisites", meta.get("conocimientosPrevios")), ("field_context", meta.get("contexto")),
                       ("field_tone", meta.get("tono")), ("field_passing", head(f"passing_{passing}")),
                       ("field_course_score", meta.get("notaAprobadoCurso"))):
        _row(sheet, [head(key), value])

    sheet = _sheet(book, head("sheet_competencies"), [(head("col_id"), 8), (head("col_text"), 90), (head("col_objectives"), 30)])
    for c in ordered(graph["competencias"]):
        served = [o["id"] for o in ordered(graph["objetivos"]) if c["id"] in (o.get("competenciaIds") or [])]
        _row(sheet, [competencies[c["id"]], c["texto"], names(served, objectives)])

    sheet = _sheet(book, head("sheet_objectives"), [(head("col_id"), 10), (head("col_text"), 70), (head("col_bloom"), 14),
                                                    (head("col_competencies"), 14), (head("col_sections"), 40),
                                                    (head("col_learning"), 30), (head("col_assessment"), 30)])
    for o in sorted(graph["objetivos"], key=lambda o: (objectives[o["id"]] if objectives[o["id"]][:1] == "U" else "~", o.get("orden", 0))):
        where = [nodes[n["id"]]["titulo"] for n in graph["nodos"] if o["id"] in (n.get("objetivoIds") or [])]
        practised = [a["titulo"] for a in graph["actividadesAprendizaje"] if o["id"] in (a.get("objetivoIds") or [])]
        assessed = [a["titulo"] for a in graph["actividadesEvaluacion"] if o["id"] in (a.get("objetivoIds") or [])]
        _row(sheet, [objectives[o["id"]], o["texto"], bloom_name.get(o.get("nivelBloom"), o.get("nivelBloom")),
                     names(o.get("competenciaIds") or [], competencies), "\n".join(where), "\n".join(practised), "\n".join(assessed)])

    sheet = _sheet(book, head("sheet_structure"), [(head("col_level"), 14), (head("col_title"), 60), (head("col_hours"), 8),
                                                   (head("col_min_words"), 12), (head("col_objectives"), 22)])
    children: dict[str | None, list[dict]] = {}
    for n in graph["nodos"]:
        children.setdefault(n["parentId"], []).append(n)

    def walk(parent: str | None, depth: int) -> None:
        for n in ordered(children.get(parent, [])):
            leaf = n in leaves
            _row(sheet, [head(f"node_{n['tipo']}"), "    " * depth + n["titulo"], n.get("horas"),
                         coursemod.min_words(n.get("horas") or 0, rules) if leaf else None, names(n.get("objetivoIds") or [], objectives)],
                 bold=n["tipo"] in ("modulo", "unidad"))
            walk(n["id"], depth + 1)
    walk(None, 0)

    unit_order = {n["id"]: i for i, n in enumerate(n for top in ordered(children.get(None, []))
                                                    for n in ([top] if top["tipo"] == "unidad" else ordered(children.get(top["id"], []))))}
    def unit_of(node_id: str) -> dict:
        """The unit a node belongs to (an activity hangs from its unit or, in slxd, from one of its sections)."""
        node = nodes[node_id]
        while node.get("parentId") and node["tipo"] != "unidad":
            node = nodes[node["parentId"]]
        return node

    def section_of(node_id: str) -> str:
        node = nodes[node_id]
        while node["tipo"] == "subapartado":
            node = nodes[node["parentId"]]
        return node["titulo"] if node["tipo"] == "apartado" else ""

    by_unit = lambda items: sorted(items, key=lambda a: (unit_order.get(unit_of(a["nodeId"])["id"], 0), a.get("orden", 0)))  # noqa: E731
    sheet = _sheet(book, head("sheet_learning"), [(head("col_unit"), 30), (head("col_section"), 30), (head("col_title"), 40),
                                                  (head("col_type"), 20), (head("col_objectives"), 16), (head("col_description"), 70)])
    for a in by_unit(graph["actividadesAprendizaje"]):
        _row(sheet, [unit_of(a["nodeId"])["titulo"], section_of(a["nodeId"]), a["titulo"], kinds.get(a.get("tipo"), a.get("tipo")),
                     names(a.get("objetivoIds") or [], objectives), a.get("descripcion")])

    sheet = _sheet(book, head("sheet_assessment"), [(head("col_unit"), 30), (head("col_title"), 40), (head("col_instrument"), 18),
                                                    (head("col_objectives"), 16), (head("col_criterion"), 50),
                                                    (head("col_weight"), 8), (head("col_passing"), 10), (head("col_attempts"), 10),
                                                    (head("col_description"), 60)])
    for a in by_unit(graph["actividadesEvaluacion"]):
        _row(sheet, [unit_of(a["nodeId"])["titulo"], a["titulo"], instruments.get(a.get("instrumento"), a.get("instrumento")),
                     names(a.get("objetivoIds") or [], objectives), a.get("criterio"), a.get("peso"), a.get("notaAprobado"),
                     a.get("intentosMax"), a.get("descripcion")])

    errors, _ = misconmod.load(course)
    sheet = _sheet(book, head("sheet_misconceptions"), [(head("col_objective"), 12), (head("col_id"), 22), (head("col_text"), 100)])
    for eid, (objective, text) in (errors.by_id.items() if errors else ()):
        _row(sheet, [objective, eid, text])
    if not errors or not errors.by_id:
        _row(sheet, ["", "", head("no_misconceptions")])

    sheet = _sheet(book, head("sheet_validation"), [(head("col_severity"), 14), (head("col_element"), 40), (head("col_check"), 30),
                                                    (head("col_message"), 90)])
    unit_word = coursemod.tokens(course)["unit"]

    def node_name(n: dict) -> str:
        up = n
        while up.get("parentId") and up["tipo"] != "unidad":
            up = nodes[up["parentId"]]
        unit = f"{unit_word} {unit_order[up['id']] + 1}" if up["id"] in unit_order else ""
        return n["titulo"] if up is n or not unit else f"{unit} · {n['titulo']}"

    element = {**{n["id"]: node_name(n) for n in graph["nodos"]}, **objectives, **competencies,
               **{a["id"]: a["titulo"] for a in graph["actividadesAprendizaje"] + graph["actividadesEvaluacion"]}}
    rank = {"error": 0, "warning": 1, "expected": 2}
    shown = [("expected" if expected(f, section_kinds) else f.severity, f) for f in findings]
    for severity, f in sorted(shown, key=lambda x: rank[x[0]]):
        _row(sheet, [head(f"severity_{severity}"), element.get(str(f.ref), ""), f.check, f.message])
    if not findings:
        _row(sheet, ["", "", "", head("no_findings")])
    book.save(target)


TEMPLATE = """\
# Instructional design of the course, made in the project (course.yaml › design.tool: local). Keys in English, texts in
# the language of the course. `coursekit design {code}` checks it and exports design/matrix.json, design/validation.json
# and the Excel the person reviews; `coursekit approve design {code}` signs that export.
profile:
  audience: ""          # who the course is for
  level: ""             # e.g. basic, intermediate
  prerequisites: ""     # what they already know
  context: ""           # where and why they take it
  tone: ""
passing: all_tests      # all_tests: every test passed | course_score: the weighted score reaches course_passing_score
# course_passing_score: 60
competencies:           # verb + object + context
  - id: C1
    text: ""
objectives:             # the course's, in order; each one: an id (U1.1, O1...), one observable verb, its Bloom level
  - id: O1              # (remember, understand, apply, analyze, evaluate, create) and the competencies it serves
    text: ""
    bloom: apply
    competencies: [C1]
units:                  # with design.structure con_modulos: modules: [{{title, hours, general_objective, units: [...]}}]
  - title: ""
    sections:           # leaves with their hours and the objectives they develop; subsections: [...] split one
      - title: ""
        hours: 0.1
        objectives: []
    learning_activities:   # type: self_assessed, classification, interactive_decision, branching_scenario,
      - title: ""          # hotspot_analysis, guided_simulation, template_practice
        type: self_assessed
        description: ""
        objectives: [O1]
    assessment_activities: # instrument (highest Bloom level it evidences): quiz (apply), case_study (evaluate),
      - title: ""          # portfolio (create), rubric (create); weight, passing_score, attempts (0 = unlimited)
        instrument: quiz
        criterion: ""
        description: ""
        objectives: [O1]
"""


def template(code: str) -> str:
    return TEMPLATE.format(code=code)
