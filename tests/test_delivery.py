import json
import re
import zipfile

import pytest
import yaml
from openpyxl import load_workbook

from coursekit import publish
from coursekit.cli import main
from tests.test_assemble import course  # noqa: F401 — the approved sample course fixture


def set_status(course_dir, status):
    path = course_dir / "course.yaml"
    path.write_text(re.sub(r"^status: .*$", f"status: {status}", path.read_text(encoding="utf-8"), count=1, flags=re.M), encoding="utf-8")


def scorm(path, manifest=True):
    with zipfile.ZipFile(path, "w") as zf:
        if manifest:
            zf.writestr("imsmanifest.xml", "<manifest/>")
        zf.writestr("index.html", "<html/>")


def test_delivery_name_and_add(course, capsys):  # noqa: F811
    assert main(["delivery", "name", "PWD", "--unit", "1", "--version", "1.0"]) == 0
    name = capsys.readouterr().out.strip()
    assert name == "PWD-U01-v1.0-scorm12.zip"
    folder = course / "delivery"
    folder.mkdir()
    scorm(folder / "bad.zip", manifest=False)
    assert main(["delivery", "add", "PWD", "--unit", "1", "--version", "1.0", "--file", str(folder / "bad.zip")]) == 1
    assert "imsmanifest.xml is not at the root" in capsys.readouterr().err
    scorm(folder / name)
    set_status(course, "assembly")
    assert main(["delivery", "add", "PWD", "--unit", "1", "--version", "1.0", "--file", str(folder / name), "--job", "j1"]) == 0
    assert "course delivered" in capsys.readouterr().out
    data = yaml.safe_load((course / "course.yaml").read_text(encoding="utf-8"))
    assert data["status"] == "delivered"
    assert data["deliveries"][0]["by"] == "Ana <ana@example.com>"
    assert data["deliveries"][0]["export_job"] == "j1"


def test_publish_to_a_folder_and_catalog(course, tmp_path, monkeypatch, capsys):  # noqa: F811
    root = course.parent.parent
    mirror = tmp_path / "mirror"
    mirror.mkdir()
    text = (root / "project.yaml").read_text(encoding="utf-8")
    text = text.replace("provider: none", "provider: sharepoint").replace(
        '  url: ""', '  url: "https://acme.sharepoint.com/sites/X/Docs/Forms/AllItems.aspx?id=%2Fsites%2FX%2FDocs%2FCursos"'
    )
    (root / "project.yaml").write_text(text, encoding="utf-8")
    monkeypatch.setenv("MIRROR_DIR", str(mirror))
    assert main(["publish", "--check"]) == 0
    assert "course folder links like https://acme.sharepoint.com" in capsys.readouterr().out
    assert main(["publish"]) == 0
    out = capsys.readouterr().out
    assert "published PWD" in out
    assert (mirror / "courses" / "PWD" / "course.yaml").exists()
    assert (mirror / "courses" / "PWD" / "content" / "unit-01" / "content.md").exists()
    ws = load_workbook(mirror / "catalogo-cursos.xlsx")["Cursos"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, [c.value for c in ws[2]], strict=True))
    assert row["Código"] == "PWD" and row["Estado"] == "DI aprobado"
    assert row["Carpeta"].endswith("id=%2Fsites%2FX%2FDocs%2FCursos%2Fcourses%2FPWD")
    # a file removed from the course is removed from the mirror; delivery/ is kept
    (course / "delivery").mkdir()
    scorm(course / "delivery" / "p.zip")
    (course / "reviews" / "r.md").write_text("x", encoding="utf-8")
    assert main(["publish", "PWD"]) == 0
    (course / "reviews" / "r.md").unlink()
    (course / "delivery" / "p.zip").unlink()
    assert main(["publish", "PWD"]) == 0
    assert not (mirror / "courses" / "PWD" / "reviews" / "r.md").exists()
    assert (mirror / "courses" / "PWD" / "delivery" / "p.zip").exists()


def test_publish_without_mirror(course, capsys):  # noqa: F811
    assert main(["publish", "--only-if-configured"]) == 0
    assert capsys.readouterr().out == ""
    assert main(["publish"]) == 0  # provider none is a choice: nothing to publish, not an error
    assert "nothing to publish" in capsys.readouterr().out


@pytest.mark.parametrize(
    ("provider", "url", "expected"),
    [
        ("nextcloud", "https://cloud.acme.org/apps/files/?dir=/Cursos", "https://cloud.acme.org/apps/files/?dir=%2FCursos%2Fcourses%2FABC"),
        ("folder", "https://nas.acme.org/share/{code}", "https://nas.acme.org/share/ABC"),
        ("google-drive", "https://drive.google.com/drive/folders/xyz", "https://drive.google.com/drive/folders/xyz"),
        ("sharepoint", "https://acme.sharepoint.com/:f:/s/sharing-token", None),
    ],
)
def test_folder_links(provider, url, expected):
    link = publish.folder_link(provider, url)
    assert (link("ABC") if link else None) == expected


def test_english_catalog(tmp_path, monkeypatch):
    monkeypatch.delenv("COURSEKIT_PROJECT", raising=False)
    root = tmp_path / "en"
    assert main(["init", str(root), "--yes", "--no-git", "--language", "en"]) == 0
    monkeypatch.chdir(root)
    assert main(["new", "Strong passwords", "1", "--code", "EN1"]) == 0
    assert main(["catalog"]) == 0
    ws = load_workbook(root / "courses" / "course-catalog.xlsx")["Courses"]
    assert [c.value for c in ws[1]][:3] == ["Code", "Title", "Status"]
    assert [c.value for c in ws[2]][2] == "Instructional design (proposal)"


def test_theme_tokens_and_check(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("COURSEKIT_PROJECT", raising=False)
    root = tmp_path / "t"
    assert main(["init", str(root), "--yes", "--no-git"]) == 0
    monkeypatch.chdir(root)
    assert main(["theme", "tokens"]) == 1
    assert "missing theme/tokens.json" in capsys.readouterr().err
    tokens = {"color": {"accent": "#FF0000", "on-accent": "#FFFFFF", "text": "#000000", "background": "#FFFFFF"},
              "font": {"family": "system-ui", "size-base": "16px", "line-height": "1.6"}, "radius": "4px",
              "space": {"sm": "8px", "md": "16px"}}
    (root / "theme" / "tokens.json").write_text(json.dumps(tokens), encoding="utf-8")
    assert main(["theme", "tokens"]) == 0
    css = (root / "theme" / "tokens.css").read_text(encoding="utf-8")
    assert "--color-accent: #FF0000;" in css and "--space-md: 16px;" in css and "@import" not in css
    capsys.readouterr()
    assert main(["theme", "check"]) == 1
    out = capsys.readouterr().out
    assert "FAIL on-accent on accent (button text on the accent): 4.00:1" in out
    tokens["color"]["accent"] = "#CC0000"
    (root / "theme" / "tokens.json").write_text(json.dumps(tokens), encoding="utf-8")
    assert main(["theme", "check"]) == 0


class Response:
    def __init__(self, body: bytes):
        self.body = body

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def read(self, size=-1):
        chunk, self.body = self.body[:size], self.body[size:]
        return chunk


def test_delivery_download_saves_and_checks_the_package(course, tmp_path, monkeypatch, capsys):  # noqa: F811
    good, bad = tmp_path / "good.zip", tmp_path / "bad.zip"
    scorm(good)
    scorm(bad, manifest=False)
    served = {"body": good.read_bytes()}
    agents = []
    monkeypatch.setattr("urllib.request.urlopen",
                        lambda req, timeout=0: agents.append(req.get_header("User-agent")) or Response(served["body"]))
    set_status(course, "assembly")
    args = ["delivery", "download", "PWD", "--unit", "1", "--version", "1.0"]
    assert main(args) == 2 and "needs --url" in capsys.readouterr().err
    assert main([*args, "--url", "http://creator.example/x"]) == 1 and "https://" in capsys.readouterr().err
    served["body"] = bad.read_bytes()
    assert main([*args, "--url", "https://creator.example/x"]) == 1
    assert not list((course / "delivery").iterdir())  # an invalid package leaves nothing behind
    served["body"] = good.read_bytes()
    assert main([*args, "--url", "https://creator.example/x"]) == 0
    path = capsys.readouterr().out.strip().splitlines()[-1]
    assert path.endswith("delivery/PWD-U01-v1.0-scorm12.zip")
    assert main(["delivery", "add", "PWD", "--unit", "1", "--version", "1.0", "--file", path]) == 0
    assert agents and all(a and a.startswith("coursekit/") for a in agents)  # Cloudflare refuses the default one of Python
