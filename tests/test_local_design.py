"""The instructional design made in the project (course.yaml › design.tool: local), without slxd: design/design.yaml,
checked and exported by `coursekit design`, then synced and signed like a design made in slxd."""

import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from coursekit import course as coursemod
from coursekit import localdesign
from coursekit.cli import main
from coursekit.project import find_project
from tests.test_course_flow import ASSESSMENT, CONTENT, write_course_rules

DESIGN = """\
profile:
  audience: Personal de oficina sin perfil técnico
  level: Inicial
  context: Migración de las herramientas de ACME a la nube
competencies:
  - id: C1
    text: Proteger sus cuentas de trabajo en el día a día
objectives:
  - id: O1
    text: Aplicar los criterios de una contraseña robusta a una cuenta nueva
    bloom: apply
    competencies: [C1]
  - id: O2
    text: Analizar por qué una contraseña es débil
    bloom: analyze
    competencies: [C1]
units:
  - title: Contraseñas que protegen
    hours: 1
    sections:
      - {title: Introducción y objetivos, hours: 0.2}
      - {title: Cómo es una contraseña robusta, hours: 0.6, objectives: [O1, O2]}
      - {title: Resumen, hours: 0.2}
    learning_activities:
      - {title: Clasifica contraseñas, type: classification, objectives: [O1, O2]}
    assessment_activities:
      - {title: Test de la unidad, instrument: case_study, criterion: Acierta la mitad, weight: 1, passing_score: 50,
         attempts: 2, objectives: [O1, O2]}
"""


@pytest.fixture
def project(tmp_path, monkeypatch):
    for key in ("COURSEKIT_PROJECT", "COURSEKIT_USER_NAME", "COURSEKIT_USER_EMAIL"):
        monkeypatch.delenv(key, raising=False)
    root = tmp_path / "demo"
    assert main(["init", str(root), "--yes", "--no-git", "--client", "Demo"]) == 0
    monkeypatch.chdir(root)
    (root / ".env").write_text("COURSEKIT_USER_NAME=Ana Pérez\nCOURSEKIT_USER_EMAIL=ana@example.com\n", encoding="utf-8")
    return root


def new_local(project: Path, design: str = DESIGN) -> Path:
    assert main(["new", "Contraseñas seguras", "1", "--code", "pwd101", "--design-tool", "local"]) == 0
    course_dir = project / "courses" / "PWD101"
    write_course_rules(course_dir)
    (course_dir / "design" / "design.yaml").write_text(design, encoding="utf-8")
    return course_dir


def findings(project: Path, design: str, **course_design) -> list[localdesign.Finding]:
    course_dir = project / "courses" / "PWD101"
    if not course_dir.exists():
        new_local(project, design)
    (course_dir / "design" / "design.yaml").write_text(design, encoding="utf-8")
    course = coursemod.load(find_project(project), "PWD101")
    course["design"].update(course_design)
    return localdesign.build(course)[1]


def test_a_local_course_designs_exports_syncs_and_signs_without_slxd(project, capsys):
    assert main(["new", "Contraseñas seguras", "1", "--code", "pwd101", "--design-tool", "local"]) == 0
    course_dir = project / "courses" / "PWD101"
    assert yaml.safe_load((course_dir / "course.yaml").read_text(encoding="utf-8"))["design"]["tool"] == "local"
    assert "objectives:" in (course_dir / "design" / "design.yaml").read_text(encoding="utf-8")  # the template to fill in
    write_course_rules(course_dir)
    (course_dir / "design" / "design.yaml").write_text(DESIGN, encoding="utf-8")
    capsys.readouterr()

    assert main(["design", "PWD101", "--check"]) == 0
    assert not (course_dir / "design" / "matrix.json").exists()  # --check writes nothing
    assert main(["design", "PWD101"]) == 0
    out = capsys.readouterr().out
    assert "without errors or warnings" in out and "design/matrix.json" in out
    graph = json.loads((course_dir / "design" / "matrix.json").read_text(encoding="utf-8"))
    assert graph["matrix"]["origen"] == "local" and graph["matrix"]["palabrasPorHora"] == 10 * 20
    assert [o["nivelBloom"] for o in graph["objetivos"]] == ["aplicar", "analizar"]
    assert graph["actividadesEvaluacion"][0]["instrumento"] == "caso_practico"
    assert json.loads((course_dir / "design" / "validation.json").read_text(encoding="utf-8")) == {"hallazgos": []}
    book = load_workbook(course_dir / "design" / "PWD101-instructional-design.xlsx")
    assert book.sheetnames[:4] == ["Profile", "Competencies", "Objectives", "Structure"]
    assert [c.value for c in book["Structure"][4]][:2] == ["Section", "    Cómo es una contraseña robusta"]

    # the same export twice gives the same matrix (stable ids): a signature does not break by exporting again
    first = (course_dir / "design" / "matrix.json").read_bytes()
    assert main(["design", "PWD101"]) == 0
    assert (course_dir / "design" / "matrix.json").read_bytes() == first

    assert main(["sync", "PWD101", "--check"]) == 0
    assert main(["approve", "design", "PWD101", "--yes", "--no-commit"]) == 0
    data = yaml.safe_load((course_dir / "course.yaml").read_text(encoding="utf-8"))
    unit = data["units"][0]
    assert [s["kind"] for s in unit["sections"]] == ["intro", "content", "summary"]
    assert [(o["id"], o["bloom"]) for o in unit["objectives"]] == [("U1.1", "aplicar"), ("U1.2", "analizar")]
    assert unit["activities"]["assessment"][0]["passing_score"] == 50

    # the rest of the flow does not care where the design was made
    (course_dir / "content" / "unit-01" / "content.md").write_text(CONTENT, encoding="utf-8")
    (course_dir / "content" / "unit-01" / "assessment.md").write_text(ASSESSMENT, encoding="utf-8")
    capsys.readouterr()
    assert main(["verify", "PWD101"]) == 0, capsys.readouterr().out

    # design.yaml edited and not exported: the signature no longer holds, and a new one is refused until it is exported
    path = course_dir / "design" / "design.yaml"
    path.write_text(path.read_text(encoding="utf-8").replace("Resumen, hours: 0.2", "Resumen, hours: 0.3"), encoding="utf-8")
    course = coursemod.load(find_project(project), "PWD101")
    assert "changed since the last export" in coursemod.approval_problem(course)
    capsys.readouterr()
    assert main(["approve", "design", "PWD101", "--yes", "--no-commit"]) == 1
    assert "coursekit design PWD101" in capsys.readouterr().err


def test_a_course_designed_in_slxd_gets_the_same_review_excel(project, capsys):
    assert main(["new", "Contraseñas seguras", "1", "--code", "pwd101"]) == 0
    design = project / "courses" / "PWD101" / "design"
    assert not (design / "design.yaml").exists()
    assert main(["design", "PWD101"]) == 1  # the design agent has not saved the matrix yet
    assert "does not exist" in capsys.readouterr().err
    matrix = json.loads((Path(__file__).parent / "fixtures" / "matrix.json").read_text(encoding="utf-8"))
    (design / "matrix.json").write_text(json.dumps(matrix), encoding="utf-8")
    intro, content = (n["id"] for n in matrix["nodos"][1:3])
    finding = lambda node: {"check": "cobertura.contenidoHuerfano", "severidad": "aviso", "mensaje": "No objective.",  # noqa: E731
                            "ref": {"tipo": "nodo", "id": node}}
    (design / "validation.json").write_text(json.dumps({"hallazgos": [finding(intro), finding(content)]}), encoding="utf-8")
    assert main(["design", "PWD101", "--check"]) == 0
    out = capsys.readouterr().out
    assert out.count("No objective.") == 1  # the introduction develops no objective by design: not reported
    assert not (design / "PWD101-instructional-design.xlsx").exists()  # --check writes nothing
    assert main(["design", "PWD101"]) == 0
    book = load_workbook(design / "PWD101-instructional-design.xlsx")
    assert "Competencies" in book.sheetnames and "Common errors" in book.sheetnames
    objectives = [r for r in book["Objectives"].iter_rows(values_only=True)][1:]
    assert [r[0] for r in objectives] == ["U1.1", "U1.2"]  # the ids course.yaml and misconceptions.yaml use
    validation = [r for r in book["Validation"].iter_rows(values_only=True)][1:]
    assert [(r[0], r[1]) for r in validation] == [("Warning", "Unidad 1 · Cómo es una contraseña robusta"),
                                                   ("Expected", "Unidad 1 · Introducción y objetivos")]


def test_the_checks_of_a_design_matrix(project):
    broken = (DESIGN.replace("instrument: case_study", "instrument: quiz")  # a quiz does not reach analyze
              .replace("text: Aplicar los criterios", "text: Comprender los criterios")
              .replace("text: Analizar por qué", "text: Evaluar por qué")
              .replace("{title: Resumen, hours: 0.2}", "{title: Extra, hours: 0.1}\n      - {title: Resumen, hours: 0.2}")
              .replace("objectives: [O1, O2]}\n    assessment", "objectives: [O2]}\n    assessment"))
    found = {f.check: f for f in findings(project, broken)}
    assert found["rigor.instrument_ceiling"].severity == "error"
    assert "level analyze and its evidence is quiz" in found["rigor.instrument_ceiling"].message
    assert found["verbs.vague"].message == "objective O1 starts with 'comprender', a verb that cannot be observed"
    assert "'evaluar', a verb of level evaluate, but declares analyze" in found["verbs.bloom_mismatch"].message
    assert found["hours.children"].message == "'Contraseñas que protegen' declares 1 h and its parts add up to 1.1 h"
    assert "Extra" in found["coverage.orphan_content"].message  # the introduction and the summary are not orphans
    assert found["arc.assessed_without_practice"].message.startswith("objective O1 is assessed without")

    found = {f.check for f in findings(project, DESIGN.replace("objectives: [O1, O2]}\n", "objectives: [O1]}\n", 1)
                                       .replace("C1]\n  - id: O2", "C1]\n  - id: O3\n    text: Detectar una contraseña reutilizada\n"
                                                "    bloom: analyze\n  - id: O2"))}
    assert {"coverage.objective_without_content", "evidence.objective_without_assessment",
            "coverage.objective_without_competency"} <= found

    found = [f.message for f in findings(project, DESIGN.replace("hours: 0.6, objectives: [O1, O2]", "hours: 0.6, objectives: [O1, O9]")
                                         .replace("bloom: apply", "bloom: aply"))]
    assert "'Cómo es una contraseña robusta' names an objective that does not exist: 'O9'" in found
    assert any("'aply' is not valid" in m for m in found)

    found = {f.check for f in findings(project, DESIGN, structure="con_modulos")}
    assert "structure.invalid" in found and "structure.empty" in found


def test_modules_and_subsections(project):
    new_local(project)
    course = coursemod.load(find_project(project), "PWD101")
    course["design"]["structure"] = "con_modulos"
    data = yaml.safe_load(DESIGN)
    data["modules"] = [{"title": "Bloque 1", "units": data.pop("units")}]
    unit = data["modules"][0]["units"][0]
    del unit["hours"]
    unit["sections"][1] = {"title": "Cómo es una contraseña robusta", "subsections": [
        {"title": "Longitud", "hours": 0.3, "objectives": ["O1"]}, {"title": "Señales de debilidad", "hours": 0.3, "objectives": ["O2"]}]}
    builder = localdesign.Builder(course, data)
    graph = builder.build()
    assert not builder.findings, builder.findings
    assert [n["tipo"] for n in graph["nodos"]] == ["modulo", "unidad", "apartado", "apartado", "subapartado", "subapartado", "apartado"]
    assert graph["nodos"][0]["horas"] == 1 and graph["nodos"][3]["horas"] == 0.6  # taken from their parts


def test_a_learning_activity_may_hang_from_a_section_of_its_unit(project):
    from coursekit import sync as syncmod

    assert main(["new", "Contraseñas seguras", "1", "--code", "pwd101"]) == 0
    matrix = json.loads((Path(__file__).parent / "fixtures" / "matrix.json").read_text(encoding="utf-8"))
    matrix["actividadesAprendizaje"] = [
        {"id": "aa-1", "nodeId": "s-2", "titulo": "Arregla la contraseña", "tipo": "practica_plantilla", "orden": 0, "objetivoIds": []},
        {"id": "aa-2", "nodeId": "u-1", "titulo": "Repaso", "tipo": "autoevaluable", "orden": 1, "objetivoIds": []}]
    design = project / "courses" / "PWD101" / "design"
    (design / "matrix.json").write_text(json.dumps(matrix), encoding="utf-8")
    (design / "validation.json").write_text('{"hallazgos": []}', encoding="utf-8")
    course = coursemod.load(find_project(project), "PWD101")
    learning = syncmod.sync(course, check_only=True).units[0]["activities"]["learning"]
    assert [(a["title"], a.get("section")) for a in learning] == [("Arregla la contraseña", 2), ("Repaso", None)]
    assert main(["design", "PWD101"]) == 0
    rows = list(load_workbook(design / "PWD101-instructional-design.xlsx")["Learning activities"].iter_rows(values_only=True))[1:]
    assert [(r[0], r[1], r[2]) for r in rows] == [("Contraseñas que protegen", "Cómo es una contraseña robusta", "Arregla la contraseña"),
                                                   ("Contraseñas que protegen", None, "Repaso")]  # an empty cell


def test_the_key_of_an_objective_in_slxd_must_be_the_id_coursekit_gives(project, capsys):
    assert main(["new", "Contraseñas seguras", "1", "--code", "pwd101"]) == 0
    matrix = json.loads((Path(__file__).parent / "fixtures" / "matrix.json").read_text(encoding="utf-8"))
    matrix["objetivos"][0]["clave"] = "U1.1"
    matrix["objetivos"][1]["clave"] = "RA2"
    (project / "courses" / "PWD101" / "design" / "matrix.json").write_text(json.dumps(matrix), encoding="utf-8")
    (project / "courses" / "PWD101" / "design" / "validation.json").write_text('{"hallazgos": []}', encoding="utf-8")
    capsys.readouterr()
    assert main(["design", "PWD101", "--check"]) == 0
    out = capsys.readouterr().out
    assert out.count("has the key") == 1 and "has the key RA2 in slxd, but coursekit calls it U1.2" in out


def test_a_validation_older_than_the_matrix_is_refused(project, capsys):
    import os

    assert main(["new", "Contraseñas seguras", "1", "--code", "pwd101"]) == 0
    design = project / "courses" / "PWD101" / "design"
    (design / "matrix.json").write_text((Path(__file__).parent / "fixtures" / "matrix.json").read_text(encoding="utf-8"),
                                        encoding="utf-8")
    assert main(["design", "PWD101"]) == 1  # never validated
    assert "validation.json is missing" in capsys.readouterr().err
    (design / "validation.json").write_text('{"hallazgos": []}', encoding="utf-8")
    later = (design / "validation.json").stat().st_mtime + 60
    os.utime(design / "matrix.json", (later, later))  # the matrix changed after it was validated
    assert main(["design", "PWD101"]) == 1
    assert "older than design/matrix.json" in capsys.readouterr().err
    assert main(["approve", "design", "PWD101", "--yes", "--no-commit"]) == 1
    assert "older than design/matrix.json" in capsys.readouterr().err
