"""Hold and resume, the client's review, the delivery gate and the signature that follows the AI review."""

import pytest
import yaml
from openpyxl import load_workbook

from coursekit import states
from coursekit.cli import main
from tests.test_assemble import course  # noqa: F401 — the approved sample course fixture
from tests.test_delivery import scorm, set_status


def data(course_dir) -> dict:
    return yaml.safe_load((course_dir / "course.yaml").read_text(encoding="utf-8"))


def link_unit(capsys):
    assert main(["assemble", "link", "PWD", "--unit", "1", "--preview", "https://creator.example/p/1", "--review",
                 "https://creator.example/r/1"]) == 0
    assert "review link https://creator.example/r/1" in capsys.readouterr().out


# ── hold and resume ──────────────────────────────────────────────────────────────────────────


def test_hold_blocks_the_work_and_resume_gives_the_status_back(course, capsys):  # noqa: F811
    assert main(["hold", "PWD", "--reason", "waiting for the client's material"]) == 0
    assert "on hold (it was in 'design_approved')" in capsys.readouterr().out
    info = data(course)
    assert info["status"] == "on_hold"
    assert info["hold"]["previous"] == "design_approved" and info["hold"]["reason"] == "waiting for the client's material"
    assert info["history"][-1]["status"] == "on_hold"

    assert main(["hold", "PWD", "--reason", "again"]) == 1
    assert "already on hold" in capsys.readouterr().err
    for command in (["write", "PWD", "1"], ["review", "PWD", "1"], ["approve", "design", "PWD", "--yes"], ["reviewed", "PWD", "1"],
                    ["assemble", "plan", "PWD", "--unit", "1"], ["sync", "PWD"], ["client", "PWD", "send", "--where", "https://x"]):
        assert main(command) == 1, command
        assert "has been on hold since" in capsys.readouterr().err, command
    assert main(["verify", "PWD", "--no-update"]) in (0, 1)  # reading is still allowed
    assert "has been on hold" not in capsys.readouterr().err
    assert main(["status", "PWD"]) == 0
    out = capsys.readouterr().out
    assert "On hold since" in out and "waiting for the client's material" in out and "coursekit resume PWD" in out

    assert main(["resume", "PWD"]) == 0
    assert "'on_hold'" not in capsys.readouterr().out
    info = data(course)
    assert info["status"] == "design_approved" and "hold" not in info
    assert main(["resume", "PWD"]) == 1
    assert "is not on hold" in capsys.readouterr().err


def test_resume_derives_the_status_while_the_course_is_being_written(course):  # noqa: F811
    set_status(course, "writing")
    states.hold(course, "pause")
    assert data(course)["status"] == "on_hold"
    previous, new = states.resume(course)
    assert previous == "writing"
    assert new == "design_approved"  # no unit has started: the status the units imply


def test_hold_needs_a_reason(course, capsys):  # noqa: F811
    with pytest.raises(SystemExit):
        main(["hold", "PWD"])
    assert "--reason" in capsys.readouterr().err


# ── the client's review ──────────────────────────────────────────────────────────────────────


def test_client_review_rounds(course, capsys):  # noqa: F811
    assert main(["client", "PWD", "send", "--to", "Eva"]) == 1
    assert "starts with the course assembled" in capsys.readouterr().err
    set_status(course, "assembly")
    assert main(["client", "PWD", "send", "--to", "Eva"]) == 1
    assert "no review links" in capsys.readouterr().err
    link_unit(capsys)

    assert main(["client", "PWD", "send", "--to", "Eva"]) == 0
    assert "round 1 opened" in capsys.readouterr().out
    info = data(course)
    assert info["status"] == "client_review"
    assert info["client_review"][0]["links"] == {"1": "https://creator.example/r/1"}
    assert info["client_review"][0]["to"] == "Eva" and info["client_review"][0]["outcome"] is None
    assert main(["client", "PWD", "send"]) == 1
    assert "round 1 is still open" in capsys.readouterr().err
    assert main(["status", "PWD"]) == 0
    assert "Client review, round 1: open · Eva" in capsys.readouterr().out

    assert main(["client", "PWD", "changes", "--note", "shorter intro"]) == 0
    assert "closed with changes" in capsys.readouterr().out
    info = data(course)
    assert info["status"] == "assembly"
    assert info["client_review"][0]["outcome"] == "changes" and info["client_review"][0]["note"] == "shorter intro"
    assert main(["client", "PWD", "changes"]) == 1
    assert "no round open" in capsys.readouterr().err

    assert main(["client", "PWD", "send", "--where", "https://review.example/all"]) == 0
    assert data(course)["client_review"][1]["round"] == 2
    capsys.readouterr()
    assert main(["client", "PWD", "approve"]) == 2
    assert "--by" in capsys.readouterr().err
    assert main(["client", "PWD", "approve", "--by", "Eva Client", "--note", "all good"]) == 0
    assert "approved by Eva Client" in capsys.readouterr().out
    info = data(course)
    assert info["client_review"][1]["outcome"] == "approved" and info["client_review"][1]["by"] == "Eva Client"
    assert info["status"] == "client_review"
    assert [h["status"] for h in info["history"]][-3:] == ["assembly", "client_review", "client_review"]
    assert main(["status", "PWD"]) == 0
    out = capsys.readouterr().out
    assert "Client review, round 2: approved" in out
    assert "deliver with /deliver PWD" in out and "wait for the client" not in out  # the round is closed


def test_the_delivery_gate_follows_the_project_setting(course, capsys):  # noqa: F811
    set_status(course, "assembly")
    assert main(["delivery", "check", "PWD"]) == 0
    assert "it can be delivered" in capsys.readouterr().out

    (course.parent.parent / "config").mkdir(exist_ok=True)
    (course.parent.parent / "config" / "delivery.yaml").write_text("client_review:\n  required: true\n", encoding="utf-8")
    assert main(["delivery", "check", "PWD"]) == 1
    assert "needs the client's review" in capsys.readouterr().err
    assert main(["delivery", "name", "PWD", "--unit", "1", "--version", "1.0"]) == 1
    capsys.readouterr()

    link_unit(capsys)
    assert main(["client", "PWD", "send", "--to", "Eva"]) == 0
    assert main(["delivery", "check", "PWD"]) == 1
    assert "waiting for the client's answer" in capsys.readouterr().err
    assert main(["client", "PWD", "changes"]) == 0
    assert main(["delivery", "check", "PWD"]) == 1
    assert "asked for changes" in capsys.readouterr().err
    assert main(["client", "PWD", "send", "--to", "Eva"]) == 0
    assert main(["client", "PWD", "approve", "--by", "Eva"]) == 0
    capsys.readouterr()
    assert main(["delivery", "check", "PWD"]) == 0

    folder = course / "delivery"
    folder.mkdir()
    name = "PWD-U01-v1.0-scorm12.zip"
    scorm(folder / name)
    assert main(["delivery", "add", "PWD", "--unit", "1", "--version", "1.0", "--file", str(folder / name)]) == 0
    assert data(course)["status"] == "delivered"


def test_skipping_the_client_review_needs_a_reason_and_is_recorded(course, capsys):  # noqa: F811
    set_status(course, "assembly")
    (course.parent.parent / "config").mkdir(exist_ok=True)
    (course.parent.parent / "config" / "delivery.yaml").write_text("client_review:\n  required: true\n", encoding="utf-8")
    assert main(["client", "PWD", "skip"]) == 2
    assert "--reason" in capsys.readouterr().err
    assert main(["client", "PWD", "skip", "--reason", "internal course"]) == 0
    info = data(course)
    assert info["client_review"][0]["outcome"] == "skipped" and info["client_review"][0]["note"] == "internal course"
    assert info["status"] == "assembly"
    assert main(["delivery", "check", "PWD"]) == 0


def test_a_course_can_override_the_project_setting(course, capsys):  # noqa: F811
    set_status(course, "assembly")
    (course.parent.parent / "config").mkdir(exist_ok=True)
    (course.parent.parent / "config" / "delivery.yaml").write_text("client_review:\n  required: true\n", encoding="utf-8")
    path = course / "course.yaml"
    path.write_text(path.read_text(encoding="utf-8") + "delivery:\n  client_review:\n    required: false\n", encoding="utf-8")
    assert main(["delivery", "check", "PWD"]) == 0


def test_a_package_recorded_outside_a_delivery_status_does_not_deliver(course, capsys):  # noqa: F811
    folder = course / "delivery"
    folder.mkdir()
    name = "PWD-U01-v1.0-scorm12.zip"
    scorm(folder / name)
    assert main(["delivery", "add", "PWD", "--unit", "1", "--version", "1.0", "--file", str(folder / name)]) == 0
    assert "not marked as delivered" in capsys.readouterr().out
    assert data(course)["status"] == "design_approved"


# ── links per unit ───────────────────────────────────────────────────────────────────────────


def test_unit_links_survive_a_sync_and_reach_the_catalog(course, capsys):  # noqa: F811
    link_unit(capsys)
    assert data(course)["units"][0]["links"]["preview"] == "https://creator.example/p/1"
    assert main(["sync", "PWD"]) == 0
    assert data(course)["units"][0]["links"]["review"] == "https://creator.example/r/1"
    out = course.parent.parent / "cat.xlsx"
    assert main(["catalog", "--output", str(out)]) == 0
    sheet = load_workbook(out)["Unidades"]
    headers = [c.value for c in sheet[1]]
    assert headers[-2:] == ["Enlace compartido", "Review"]
    assert [c.value for c in sheet[2]][-2:] == ["https://creator.example/p/1", "https://creator.example/r/1"]
    assert "Enlace compartido" not in [c.value for c in load_workbook(out)["Cursos"][1]]


def test_assemble_link_needs_something_to_record(course, capsys):  # noqa: F811
    assert main(["assemble", "link", "PWD", "--unit", "1"]) == 2
    assert "--content-id, --preview or --review" in capsys.readouterr().err


# ── the signature follows the AI review ──────────────────────────────────────────────────────


def test_approve_refuses_a_unit_changed_after_its_ai_review(course, capsys):  # noqa: F811
    assert main(["verify", "PWD"]) == 0
    (course / "reviews").mkdir(exist_ok=True)
    (course / "reviews" / "unit-01-ai-review.md").write_text("# Review\n\nNo findings.\n", encoding="utf-8")
    assert main(["reviewed", "PWD", "1"]) == 0
    content = course / "content" / "unit-01" / "content.md"
    content.write_text(content.read_text(encoding="utf-8").replace("Más larga es mejor.", "Cuanto más larga, mejor."), encoding="utf-8")
    capsys.readouterr()
    assert main(["approve", "content", "PWD", "--unit", "1", "--yes", "--no-commit"]) == 1
    err = capsys.readouterr().err
    assert "changed after its AI review (section 2)" in err and "--force" in err
    assert data(course)["units"][0]["status"] == "reviewed"
    assert main(["approve", "content", "PWD", "--unit", "1", "--yes", "--no-commit", "--force"]) == 0
    assert data(course)["units"][0]["status"] == "approved"


def test_approve_without_changes_does_not_need_force(course, capsys):  # noqa: F811
    assert main(["verify", "PWD"]) == 0
    (course / "reviews").mkdir(exist_ok=True)
    (course / "reviews" / "unit-01-ai-review.md").write_text("# Review\n", encoding="utf-8")
    assert main(["reviewed", "PWD", "1"]) == 0
    assert main(["approve", "content", "PWD", "--unit", "1", "--yes", "--no-commit"]) == 0
    assert data(course)["units"][0]["status"] == "approved"


# ── the mirror follows the course by itself ──────────────────────────────────────────────────


def with_mirror(course, tmp_path, monkeypatch):  # noqa: F811
    root = course.parent.parent
    mirror = tmp_path / "mirror"
    mirror.mkdir()
    text = (root / "project.yaml").read_text(encoding="utf-8").replace("provider: none", "provider: folder")
    (root / "project.yaml").write_text(text, encoding="utf-8")
    monkeypatch.setenv("MIRROR_DIR", str(mirror))
    return mirror


def test_the_commands_that_change_a_course_update_the_mirror_and_its_catalog(course, tmp_path, monkeypatch, capsys):  # noqa: F811
    mirror = with_mirror(course, tmp_path, monkeypatch)
    link_unit(capsys)
    assert (mirror / "courses" / "PWD" / "course.yaml").exists()
    assert (course.parent / "catalogo-cursos.xlsx").exists()  # the project keeps its own copy next to the courses
    sheet = load_workbook(mirror / "catalogo-cursos.xlsx")["Unidades"]
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    assert "https://creator.example/r/1" in rows[1] and "https://creator.example/p/1" in rows[1]

    assert main(["hold", "PWD", "--reason", "x"]) == 0
    assert "published PWD" in capsys.readouterr().out
    assert "on_hold" in (mirror / "courses" / "PWD" / "course.yaml").read_text(encoding="utf-8")


def test_reading_commands_do_not_publish(course, tmp_path, monkeypatch, capsys):  # noqa: F811
    mirror = with_mirror(course, tmp_path, monkeypatch)
    for command in (["sync", "PWD", "--check"], ["verify", "PWD", "--no-update"], ["status", "PWD"],
                    ["assemble", "plan", "PWD", "--unit", "1"]):
        main(command)
        capsys.readouterr()
    assert not (mirror / "catalogo-cursos.xlsx").exists()


def test_without_a_mirror_nothing_is_published_or_printed(course, capsys):  # noqa: F811
    link_unit(capsys)  # provider none in the sample project: no extra lines, no warning
    assert main(["hold", "PWD", "--reason", "x"]) == 0
    captured = capsys.readouterr()
    assert "published" not in captured.out and "mirror" not in captured.err
    # the catalog of the project is written anyway, in courses/ and without a word
    sheet = load_workbook(course.parent / "catalogo-cursos.xlsx")["Unidades"]
    assert "https://creator.example/r/1" in [c.value for c in sheet[2]]
    assert "wrote" not in captured.out


def test_a_mirror_that_cannot_be_written_is_a_warning_not_a_failure(course, tmp_path, monkeypatch, capsys):  # noqa: F811
    mirror = with_mirror(course, tmp_path, monkeypatch)
    mirror.rmdir()
    assert main(["hold", "PWD", "--reason", "x"]) == 0
    captured = capsys.readouterr()
    assert "the mirror folder could not be updated" in captured.err
    assert data(course)["status"] == "on_hold"


def test_the_catalog_of_the_project_is_next_to_the_courses_and_not_versioned(tmp_path, monkeypatch):
    monkeypatch.delenv("COURSEKIT_PROJECT", raising=False)
    root = tmp_path / "demo"
    assert main(["init", str(root), "--yes", "--no-git", "--language", "en"]) == 0
    assert "courses/*.xlsx" in (root / ".gitignore").read_text(encoding="utf-8").splitlines()
    monkeypatch.chdir(root)
    assert main(["new", "Strong passwords", "1", "--code", "EN1"]) == 0
    assert (root / "courses" / "course-catalog.xlsx").exists()  # `coursekit new` already updated it
    assert main(["publish"]) == 0  # no mirror: it still writes the catalog of the project
    assert [p.name for p in (root / "courses").iterdir() if p.is_file()] == ["course-catalog.xlsx"]
